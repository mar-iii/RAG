from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import openpyxl
from pypdf import PdfReader
import xlrd

from app.config import settings

SUPPORTED_EXTENSIONS = (".txt", ".md", ".pdf", ".xlsx", ".xlsm", ".xls")


@dataclass(frozen=True)
class PreparedDocument:
    filename: str
    file_type: str
    text: str


def split_text(text: str) -> list[str]:
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + settings.chunk_size, len(text))
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(text):
            break
        start = end - settings.chunk_overlap
    return chunks


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_openpyxl(content: bytes) -> str:
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None]
            if values:
                rows.append(" | ".join(values))
        if rows:
            sections.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))
    return "\n\n".join(sections)


def _extract_xls(content: bytes) -> str:
    workbook = xlrd.open_workbook(file_contents=content)
    sections: list[str] = []
    for sheet in workbook.sheets():
        rows: list[str] = []
        for row_index in range(sheet.nrows):
            values = [
                str(value).strip()
                for value in sheet.row_values(row_index)
                if value != ""
            ]
            if values:
                rows.append(" | ".join(values))
        if rows:
            sections.append(f"Sheet: {sheet.name}\n" + "\n".join(rows))
    return "\n\n".join(sections)


def prepare_document(filename: str, content: bytes) -> PreparedDocument:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(SUPPORTED_EXTENSIONS)
        raise ValueError(f"Unsupported file type. Supported types: {supported}")
    if len(content) > settings.max_upload_size_mb * 1024 * 1024:
        raise ValueError(
            f"File exceeds the {settings.max_upload_size_mb} MB upload limit"
        )

    try:
        if extension in (".txt", ".md"):
            text = content.decode("utf-8")
        elif extension == ".pdf":
            text = _extract_pdf(content)
        elif extension in (".xlsx", ".xlsm"):
            text = _extract_openpyxl(content)
        else:
            text = _extract_xls(content)
    except (UnicodeDecodeError, ValueError, OSError, xlrd.XLRDError) as exc:
        raise ValueError(f"Could not read {filename}: {exc}") from exc

    text = text.strip()
    if not text:
        raise ValueError(
            f"No extractable text found in {filename}. Scanned PDFs may require OCR."
        )
    return PreparedDocument(filename=filename, file_type=extension[1:], text=text)
